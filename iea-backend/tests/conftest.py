import io
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

SAMPLES_DIR = Path(__file__).resolve().parents[1] / "samples"


@pytest.fixture(scope="session")
def synthetic_xlsx_bytes() -> bytes:
    path = SAMPLES_DIR / "synthetic_manufacturing.xlsx"
    if not path.exists():
        pytest.skip("Run samples/generate_sample.py first")
    return path.read_bytes()


@pytest.fixture(scope="session")
def synthetic_xlsm_bytes() -> bytes:
    path = SAMPLES_DIR / "synthetic_manufacturing_macro.xlsm"
    if not path.exists():
        pytest.skip("Run samples/generate_sample_xlsm.py first")
    return path.read_bytes()


@pytest.fixture
def minimal_xlsx_bytes() -> bytes:
    wb = Workbook()
    wb.active.append(["A", "B"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def workbook_with_hidden_sheet_bytes() -> bytes:
    wb = Workbook()
    wb.active.title = "Visible"
    wb.active.append(["X"])
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["Y"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def workbook_with_formulas_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["A", "B", "Sum"])
    ws.append([1, 2, "=A2+B2"])
    ws.append([3, 4, "=A3+B3"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def workbook_with_named_range_bytes() -> bytes:
    from openpyxl.workbook.defined_name import DefinedName

    wb = Workbook()
    ws = wb.active
    ws["A1"] = 42
    wb.defined_names["MyRange"] = DefinedName("MyRange", attr_text=f"{ws.title}!$A$1")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def empty_worksheet_bytes() -> bytes:
    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def malformed_xlsx_bytes() -> bytes:
    # Valid zip magic bytes but not a real zip/xlsx structure past that.
    return b"PK\x03\x04" + b"not a real zip archive" * 5


@pytest.fixture
def not_a_zip_bytes() -> bytes:
    return b"this is just a plain text file pretending to be an xlsx"


@pytest.fixture
def zip_bomb_like_bytes() -> bytes:
    """A small zip file that expands far beyond MAX_UNCOMPRESSED_BYTES."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        chunk = b"0" * (1024 * 1024)  # 1 MB of repeated zero bytes compresses tiny
        for i in range(300):  # ~300 MB uncompressed
            zf.writestr(f"junk_{i}.bin", chunk)
    return buf.getvalue()

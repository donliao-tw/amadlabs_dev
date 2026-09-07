"""Generates a SYNTHETIC manufacturing workbook for demos and tests.

All data below is fabricated for illustration only — no real company,
supplier, or part data. Run with:

    python samples/generate_sample.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT_PATH = Path(__file__).parent / "synthetic_manufacturing.xlsx"


def build() -> Workbook:
    wb = Workbook()

    suppliers = wb.active
    suppliers.title = "Suppliers"
    suppliers.append(["SupplierID", "SupplierName", "Country", "ApprovedSince"])
    suppliers.append(["S001", "Northwind Metal Co (SYNTHETIC)", "TW", "2021-03-01"])
    suppliers.append(["S002", "Bluepeak Fasteners (SYNTHETIC)", "VN", "2022-07-15"])
    suppliers.append(["S002", "Bluepeak Fasteners (SYNTHETIC)", "VN", "2022-07-15"])  # intentional duplicate
    suppliers.append(["S003", "Rivertown Plastics (SYNTHETIC)", "TW", None])  # missing ApprovedSince

    parts = wb.create_sheet("Parts")
    parts.append(["PartNo", "Description", "SupplierID", "UnitCost", "MOQ"])
    parts.append(["P-1001", "Bracket A (SYNTHETIC)", "S001", 1.25, 500])
    parts.append(["P-1002", "Fastener M4 (SYNTHETIC)", "S002", 0.03, 10000])
    parts.append(["P-1003", "Housing Shell (SYNTHETIC)", "S003", 4.80, 200])
    parts.append(["P-1004", "Bracket A (SYNTHETIC)", "S001", 1.30, 500])  # inconsistent cost for same-looking part

    inspections = wb.create_sheet("Inspections")
    inspections.append(["InspectionID", "PartNo", "LotSize", "DefectCount", "DefectRate", "Status"])
    inspections.append(["INS-001", "P-1001", 500, 3, "=D2/C2", "PASS"])
    inspections.append(["INS-002", "P-1002", 10000, 620, "=D3/C3", "FAIL"])
    inspections.append(["INS-003", "P-1003", 200, 4, "=D4/C4", "Pass"])  # inconsistent casing vs "PASS"
    inspections.append(["INS-004", "P-1001", 500, 0, "=D5/C5", "closed"])  # different status vocabulary

    dv = DataValidation(type="list", formula1='"PASS,FAIL,Pass,closed"', allow_blank=True)
    inspections.add_data_validation(dv)
    dv.add("F2:F5")

    wb.create_sheet("_ScratchCalc").sheet_state = "hidden"
    scratch = wb["_ScratchCalc"]
    scratch["A1"] = "internal helper sheet (SYNTHETIC)"
    scratch["A2"] = "=SUM(Inspections!D2:D5)"

    inspections["G1"] = 0.05
    wb.defined_names["DefectThreshold"] = DefinedName("DefectThreshold", attr_text="Inspections!$G$1")

    return wb


if __name__ == "__main__":
    workbook = build()
    workbook.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")
